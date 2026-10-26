"""Verify a student self-assessment without touching lecturer inputs."""
import hashlib
import json
from pathlib import Path
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
source=ROOT.parent/'SE_FP_SaigonBite_MarkingSheet_Sem01_2627.xlsx'
assert hashlib.sha256(source.read_bytes()).hexdigest()=='13a7cf73629bcd98f4b8b9e348c1db255f808ec56a89e097d1b41317967daae7'
team=json.loads((ROOT/'config/team.json').read_text())
output=ROOT/f"docs/part1-proposal-req/{team['project']['team_id']}_SE_MarkingSheet_SaigonBite.xlsx"
original=openpyxl.load_workbook(source)
wb=openpyxl.load_workbook(output)
cached=openpyxl.load_workbook(output,data_only=True)
assert wb.sheetnames==original.sheetnames and len(wb.sheetnames)==5
rows=list(range(13,20))+list(range(21,28))+list(range(29,49))+list(range(50,56))
assert len(rows)==40
editable={
 'Summary':{'B4','B6','B24'}|{f'{c}{r}' for c in 'BCD' for r in range(18,21)},
 'Marking':{'E5','E6'}|{f'H{r}' for r in range(5,13)},
 'Checklist':{f'{c}{r}' for c in 'DF' for r in rows},
 }
def colour(c):
    return c.rgb[-6:] if c and c.type=='rgb' else c

for old in original:
    new=wb[old.title]
    assert str(old.merged_cells)==str(new.merged_cells), (old.title,'merges')
    assert old.freeze_panes==new.freeze_panes, (old.title,'freeze')
    assert old.protection.sheet==new.protection.sheet
    assert len(old.conditional_formatting)==len(new.conditional_formatting), (old.title,'conditional formatting')
    for row in old:
        for c in row:
            n=new[c.coordinate]
            if c.coordinate not in editable.get(old.title,set()):
                assert c.value==n.value, (old.title,c.coordinate,'template value/formula changed')
            if c.has_style and not isinstance(c,MergedCell):
                assert (c.font.name,c.font.sz,c.font.bold,colour(c.font.color))==(n.font.name,n.font.sz,n.font.bold,colour(n.font.color)), (old.title,c.coordinate,'font')
                assert c.fill.patternType==n.fill.patternType and colour(c.fill.fgColor)==colour(n.fill.fgColor), (old.title,c.coordinate,'fill')
            if c.fill.fgColor.type=='rgb' and c.fill.fgColor.rgb[-6:]=='DDF0DF':
                assert c.value==n.value, (old.title,c.coordinate,'lecturer input')
    for rule in old.data_validations.dataValidation:
        for area in rule.sqref.ranges:
            for row,col in area.cells:
                address=f'{get_column_letter(col)}{row}'
                assert any(address in r and r.type==rule.type and r.formula1==rule.formula1 and r.formula2==rule.formula2 for r in new.data_validations.dataValidation), (old.title,address,'validation')
assert wb['Summary']['B4'].value=='TeamG4' and wb['Summary']['B6'].value=='Tier 4'
for i,m in enumerate(team['members'],18):
    assert [wb['Summary'].cell(i,c).value for c in [2,3]]==[m['full_name'],m['student_id']]
    assert m['use_case_title'] in wb['Summary'].cell(i,4).value
    assert wb['Summary'].cell(i,5).value is None and wb['Summary'].cell(i,6).value is None
    reviewer=['524H0179','524H0195','524H0191'][i-18]
    assert reviewer != m['student_id'] and f'Planned reviewer: {reviewer}' in wb['Summary'].cell(i,4).value
assert 'Planned peer review:' in wb['Marking']['H5'].value
assert 'Planned peer review:' in wb['Checklist']['F14'].value
assert 'incomplete whole-project subtotal' in wb['Summary']['B24'].value
assert 'not assigned marks' in wb['Summary']['B24'].value
assert [wb['Marking'][p].value for p in ['E5','E6']]==[14.5,4.5]
assert 'not independently verified' in wb['Marking']['H6'].value
assert 'Team self-assessment' in wb['Checklist']['F19'].value
assert 'not independently verified' in wb['Checklist']['F19'].value
for row in range(5,13):
    assert wb['Marking'].cell(row,6).value is None and wb['Marking'].cell(row,9).value is None
    assert wb['Marking'].cell(row,8).value
    if row>6:assert wb['Marking'].cell(row,5).value is None
assert [cached['Marking'][a].value for a in ['C13','E13','F13']]==[100,19,0]
assert cached['Marking']['E17'].value==19
assert cached['Summary']['B7'].value==10 and cached['Summary']['B14'].value==1.9
assert [cached['Checklist'].cell(r,4).value for r in range(5,9)]==[7,2,31,0]
assert cached['Checklist']['D9'].value==0.2
for row in rows:
    assert wb['Checklist'].cell(row,4).value in ['Yes','Partial','No','N/A']
    assert wb['Checklist'].cell(row,5).value is None and wb['Checklist'].cell(row,6).value
assert all(wb['Checklist'].cell(row,4).value=='Yes' for row in range(13,20))
assert all(wb['Checklist'].cell(row,4).value=='Partial' for row in [44,54])
assert all(wb['Checklist'].cell(row,4).value=='No' for row in [46,47,48])
for s in cached:
    for row in s:
        for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
print('PASS marking sheet: original5 sheets/headers/styles/formulas/validations retained; team self-estimate19/20 (14.5+4.5), pitch Yes explicitly self-assessed/unverified; all40 statuses documented; lecturer inputs blank; native recalculation and zero formula errors.')
