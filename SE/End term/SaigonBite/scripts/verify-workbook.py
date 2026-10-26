"""Read-only content, template fidelity and cached-calculation checks."""
import hashlib
import json
from pathlib import Path
import openpyxl
from openpyxl.utils.cell import get_column_letter
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parents[1]
source = ROOT.parent / 'SE_Sem01_2627_Team_Workbook_Template.xlsx'
assert hashlib.sha256(source.read_bytes()).hexdigest() == 'e6efb91f2a497825ee9707b54cab24ac38d47054ad8bdd5bac389f228c957dab'
data = json.loads((ROOT / 'config/part1.json').read_text())
team = json.loads((ROOT / 'config/team.json').read_text())
original = openpyxl.load_workbook(source)
output = ROOT / f"docs/part1-proposal-req/{team['project']['team_id']}_SE_Team_Workbook_SaigonBite.xlsx"
wb = openpyxl.load_workbook(output)
cached = openpyxl.load_workbook(output, data_only=True)
assert wb.sheetnames == original.sheetnames and len(wb.sheetnames) == 14, wb.sheetnames

def rgb(colour):
    return colour.rgb[-6:] if colour and colour.type == 'rgb' else colour

def same_style(a, b):
    assert a.font.name == b.font.name and a.font.sz == b.font.sz, (a.coordinate, 'font')
    assert a.font.bold == b.font.bold and rgb(a.font.color) == rgb(b.font.color), (a.coordinate, 'font colour/bold')
    assert a.fill.patternType == b.fill.patternType and rgb(a.fill.fgColor) == rgb(b.fill.fgColor), (a.coordinate, 'fill')

# Every template heading and original student/instructor colour remains recognisable.
for old in original:
    new = wb[old.title]
    assert str(old.merged_cells) == str(new.merged_cells), (old.title, 'merges')
    assert old.freeze_panes == new.freeze_panes, (old.title, 'freeze panes')
    for row in [1, 2, 4]:
        for c in old[row]:
            if c.value is not None:
                assert c.value == new[c.coordinate].value, (old.title, c.coordinate, 'heading')
            if c.has_style and not isinstance(c,MergedCell): same_style(c, new[c.coordinate])
    for row in old:
        for c in row:
            if c.fill.fgColor.type == 'rgb' and c.fill.fgColor.rgb[-6:] == 'DDF0DF':
                assert c.value == new[c.coordinate].value, (old.title, c.coordinate, 'instructor cell changed')
                same_style(c,new[c.coordinate])
    for rule in old.data_validations.dataValidation:
        for area in rule.sqref.ranges:
            for row,col in area.cells:
                address = f'{get_column_letter(col)}{row}'
                assert any(address in r and r.formula1 == rule.formula1 for r in new.data_validations.dataValidation), (old.title,address,'dropdown')

# Later-phase recording sheets stay blank and retain all fixed rules/formulas.
untouched = ['Instructions','EP-BVA Planner','Test Cases','State Tests','Change Request','Feedback Log']
for name in untouched:
    for row in original[name]:
        for c in row:
            assert c.value == wb[name][c.coordinate].value, (name,c.coordinate,'unrelated content')
            if c.has_style and not isinstance(c,MergedCell): same_style(c,wb[name][c.coordinate])
    assert len(original[name].conditional_formatting) == len(wb[name].conditional_formatting), (name,'conditional formatting')

for old in original:
    for row in old:
        for c in row:
            if c.data_type != 'f': continue
            if old.title == 'Requirements': continue  # Stable IDs and relocated totals.
            if old.title == 'RTM' and c.row >= 51: continue
            assert c.value == wb[old.title][c.coordinate].value, (old.title,c.coordinate,'original formula')

req = wb['Requirements']
records = data['fr'] + data['nfr']
for row, r in enumerate(records,5):
    assert req.cell(row,1).value == r['id']
    expected = r.get('statement',f"{r.get('attribute')}: {r.get('target')}")
    assert req.cell(row,2).value == expected, r['id']
    assert req.cell(row,3).value == ('FR' if row<52 else 'NFR')
    assert req.cell(row,5).value == r['use_case']
    assert req.cell(row,6).value == r['source']
    assert req.cell(row,8).value == r['owner']
    assert req.cell(row,10).value is None and req.cell(row,11).value == 'Draft'
    if 'verification' in r: assert r['verification'] in req.cell(row,9).value
    assert cached['RTM'].cell(row,1).value == r['id'], ('RTM link',r['id'])
    assert cached['RTM'].cell(row,2).value == expected
    assert wb['RTM'].cell(row,7).value == 'Not tested'
    assert cached['RTM'].cell(row,8).value == 'OK'  # Test reference exists, not execution.
    for col in [3,4,11]:
        assert any(req.cell(row,col).coordinate in r for r in req.data_validations.dataValidation), ('extended dropdown',row,col)
    assert any(f'G{row}' in r for r in wb['RTM'].data_validations.dataValidation)
    for col in range(1,12): same_style(original['Requirements'].cell(49,col),req.cell(row,col))
assert [cached['Requirements'][p].value for p in ['C59','D59','D60']] == [47,6,'OK']
assert [cached['RTM'][p].value for p in ['C59','C60']] == [0,0]
for row,r in enumerate(data['ambiguities'],5):
    s=wb['Ambiguity Log']
    assert [s.cell(row,c).value for c in [1,5,6,7,8,9,10]] == [r['id'],r['assumption'],'Draft',None,None,r['assumption'],r['impact']]
assert [cached['Ambiguity Log'][p].value for p in ['D21','D22']] == [5,'OK']
assert wb['Team & Tier']['B4'].value == 'TeamG4'
assert wb['Team & Tier']['B5'].value == 'Tier 4' and wb['Team & Tier']['B7'].value is None
assert [cached['Team & Tier'][p].value for p in ['B6','B8','D27']] == [10,3,'OK']
assert wb['Team & Tier']['F18'].value == team['members'][0]['student_id']
assert wb['Team & Tier']['D28'].value == team['members'][0]['student_id']
reviewers = ['524H0179','524H0195','524H0191']
for i,m in enumerate(team['members']):
    assert [wb['Team & Tier'].cell(i+11,c).value for c in [2,3]] == [m['full_name'],m['student_id']]
    assert cached['Team & Tier'].cell(i+11,7).value == 'OK'
    assert wb['Contribution & Commits'].cell(i+5,4).value == 'Not verified'
    assert reviewers[i] != m['student_id']
    assert f'Planned peer reviewer: {reviewers[i]}' in wb['Team & Tier'].cell(i+11,4).value
    assert f'Planned peer reviewer: {reviewers[i]}' in wb['Contribution & Commits'].cell(i+5,6).value
    assert 'not audited' in wb['Contribution & Commits'].cell(i+5,7).value
for row in range(24,27):
    assert wb['Plan & Risks'].cell(row,6).value == 'To do'
    assert wb['Plan & Risks'].cell(row,7).value.startswith('Planned')
assert 'collaboration' in wb['AI Usage Log']['C9'].value
for row in range(37,42):
    for col in range(1,4):
        assert wb['Team & Tier'].cell(row,col).value == original['Team & Tier'].cell(row,col).value
for s in cached:
    for row in s:
        for c in row:
            assert c.data_type != 'e', (s.title,c.coordinate,c.value)
print('PASS workbook: 14 original sheets; headings, input colours, merges, panes, dropdowns, untouched future sheets; 47 FR/6 NFR/5 ambiguities; 53 linked RTM rows; original tier/ownership formulas and cached counts; no formula errors or invented instructor verdicts.')
