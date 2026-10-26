"""Build the report from the single executed notebook's exported artifacts."""
import csv
import json
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from report_layout import Report

ROOT=Path(__file__).resolve().parents[1]
(ROOT/'.build').mkdir(exist_ok=True)
META=json.loads((ROOT/'project.json').read_text())
S=json.loads((ROOT/'results/summary.json').read_text())
R=Report(ROOT,META)


def paras(*texts):
    for text in texts:R.para(text)


def signed_names():
    for text in ['Ho Chi Minh City, 20 October 2026','Authors','(Full names)']+[a['name'].title() for a in META['authors']]:
        p=R.para(text,first=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent=Cm(7.5)
        p.paragraph_format.line_spacing=1.15
        if text in ['Authors','(Full names)'] or text.startswith('Ho Chi Minh'):
            for run in p.runs:run.italic=True


def hyperlink(paragraph,url):
    import docx.opc.constants
    rid=R.doc.part.relate_to(url,docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
    element=OxmlElement('w:hyperlink');element.set(qn('r:id'),rid)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');props.append(font)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'26');props.append(size);run.append(props)
    text=OxmlElement('w:t');text.text=url;run.append(text);element.append(run);paragraph._p.append(element)


R.cover();R.cover(inner=True)
R.new_numbered_section('lowerRoman',1,size=12)
R.title('ACKNOWLEDGEMENT')
paras(
    'We thank Ho Thi Linh, Ph.D., for teaching Data Analysis and Visualization and for providing the endterm assignment requirements. The assignment brings together data acquisition, missing-value handling, Python data structures, statistical analysis and visual communication in one reproducible workflow.',
    'We acknowledge Paulo Cortez and the UCI Machine Learning Repository for making the Student Performance dataset publicly available. The accompanying data documentation makes it possible to distinguish original measurements, questionnaire categories and the derived variables used in this analysis.',
    'The notebook, report and presentation were prepared with computational and AI assistance. Public data sources and software documentation are identified, and numerical statements are tied to the executed analysis. This report is intended to support an informed presentation of the methods, results and their limitations.'
)
R.doc.add_paragraph();signed_names()
R.doc.add_page_break();R.title('DECLARATION OF AUTHORSHIP')
paras(
    'This report presents a reproducible analysis of the publicly available UCI Student Performance dataset for the Portuguese-language subject. The original data belong to the cited source and are reused under the stated Creative Commons Attribution 4.0 licence.',
    'The data-processing code, analysis and document preparation were assisted by OpenAI Codex. The student authors are responsible for checking the content, understanding the methods, and complying with course rules concerning the use and disclosure of assistance. This document does not claim that the public dataset was collected by the group or that every line of code was written without assistance.',
    'The missing values used for the preprocessing exercise were introduced deliberately and are labelled as simulated. The substantive analysis uses the validated original measurements. The report distinguishes observed associations from causal effects, and it does not invent signatures, field observations or contribution percentages.'
)
R.para('The authors remain responsible for the accuracy of the submitted content and its source attribution.',bold=True)
R.doc.add_paragraph();signed_names()
R.new_numbered_section('lowerRoman',None,size=13)
R.index_placeholder('TOC','TABLE OF CONTENT');R.doc.add_page_break()
R.index_placeholder('LOF','LIST OF FIGURES');R.doc.add_page_break()
R.index_placeholder('LOT','LIST OF TABLES')
R.new_numbered_section('decimal',1)

R.heading('CHAPTER 1. INTRODUCTION')
R.heading('1.1 Objectives and research questions',2)
paras(
    'This endterm project applies Pandas, NumPy, Seaborn and Matplotlib to a small, documented educational dataset. Its purpose is to demonstrate a complete analytical workflow: inspecting a source, simulating and treating missing data, creating features, grouping and transforming observations, visualizing distributions and relationships, and reporting conclusions supported by executed code.',
    'Four questions organize the analysis. First, how are final grades distributed across the observed records? Second, how do mean imputation and row removal change a simulated incomplete dataset? Third, which original numerical or ordinal factors show the strongest observed association with the final grade? Fourth, how do grouped summaries and visual comparisons qualify that association?',
    'The exercise is descriptive. It does not estimate an intervention, train a prediction system or make a causal claim. This distinction is especially important when earlier grades are compared with final grades: repeated assessments of related achievement can be highly correlated without revealing why performance differs.'
)
R.heading('1.2 Dataset and provenance',2)
paras(
    'The source is Student Performance by Cortez (2008), hosted by the UCI Machine Learning Repository. It contains school records and questionnaire information from two Portuguese secondary schools. Only student-por.csv, the Portuguese-language subject, is used. The mathematics file is excluded because the source documents overlap between subjects; concatenating them would not represent a simple sample of independent students.',
    f"The selected CSV contains {S['dataset']['rows']} records and {S['dataset']['columns']} original columns, with no missing cells. GP contributes {S['dataset']['school_counts']['GP']} records and MS contributes {S['dataset']['school_counts']['MS']}. G1 and G2 are earlier-period grades; G3 is the final-period outcome, measured on a theoretical scale from 0 to 20. The observed G3 range is {S['dataset']['observed_G3_range'][0]}–{S['dataset']['observed_G3_range'][1]}, which need not span the whole theoretical scale.",
    'The raw CSV is preserved unchanged. Its SHA-256 checksum, download location, retrieval timestamp, licence and subject choice are recorded in references/dataset_source.json. The original codebook is retained as references/student.txt. The file is small enough to include in the submission, enabling analysis without a network connection after dependencies have been installed.'
)
R.table(['Field family','Examples','Interpretation'],[
    ['Grades','G1, G2, G3','Assessment scores on a 0–20 scale'],
    ['Quantitative','age, absences','Recorded age and nonnegative absence count'],
    ['Ordered categories','studytime, Medu, Fedu','Codes have order; intervals are not necessarily equal'],
    ['School/context','school, address, higher','Categorical labels, not arbitrary numeric distances'],
    ['Lifestyle ratings','famrel, freetime, health','Ordered questionnaire responses'],
], '1.2.1','Selected variable families and measurement levels',[3.3,4.5,7.7])
R.heading('1.3 Reproducibility and scope',2)
paras(
    'The executed Jupyter notebook is the computational source of truth. It exports tables, figures and a structured summary used by both document builders. Report and slide builders format those artifacts; they do not implement a second statistical analysis. The manifest records the raw checksum, simulation seed, dependency versions and hashes of analytical artifacts.',
    'All four assignment parts are cross-referenced in REQUIREMENTS.md. The notebook includes the required head/tail, indexing, missingness, feature engineering, group statistics, transformation, plotting and reporting operations. Additional boundary tests check the failure-prone cases: score-band edges, constant groups and summaries with zero-count categories.'
)

R.heading('CHAPTER 2. ACQUISITION AND PREPROCESSING',new_page=True)
R.heading('2.1 Loading and structural inspection',2)
paras(
    'Pandas reads the CSV with a semicolon delimiter. Before transformation, the notebook verifies the complete column order, shape, original missing counts, numeric storage types and selected domains. It displays the first five and last five records, the dtypes of every column and the output of DataFrame.info(). These checks would expose a delimiter error, a subject substitution or an unanticipated source revision.',
    f"There are {S['dataset']['duplicate_attribute_rows']} exact duplicate full attribute rows in the selected file. Duplicate attribute combinations, if present, would not be deleted automatically: anonymized attributes are insufficient to establish that two records are the same individual. Source row positions remain available as zero-based indexes in the exported tables.",
    'Some integer-valued fields are ordinal questionnaire codes. For example, studytime values 1–4 denote less than 2, 2–5, 5–10 and more than 10 weekly study hours. Treating these values as exact hours would be incorrect. The upper category of failures is ambiguously described in the original codebook; observed codes are retained as recorded, without claiming that the upper code measures an exact number of failures.'
)
R.heading('2.2 Controlled missing-data experiment',2)
paras(
    f"The original dataset is complete. To satisfy the missing-data requirement, a seeded random generator selects {S['missing_simulation']['masked_count']} distinct row indexes and replaces only their absences values with NaN. The seed is 42 and selection is without replacement. Since 5% of 649 is 32.45, rounding produces 32 masked cells, or {S['missing_simulation']['percent']:.4f}% of records. The report therefore describes the experiment as approximately 5%, not exactly 5%.",
    'The selection mechanism is missing completely at random by construction. This says nothing about how absence information would go missing in a real school system. The mask is exported so that both treatment branches can be verified against exactly the same incomplete observations. All other columns and all unmasked absences values are asserted to remain equal to the original data.'
)
R.equation('p','32 × 100','649')
R.heading('2.3 Mean imputation and row removal',2)
paras(
    f"Mean imputation replaces missing absences with the mean of the remaining observed values, {S['missing_simulation']['observed_mean']:.6f}. The mean is calculated after masking; the concealed original values are not used to fit the replacement. The alternative branch applies dropna() directly to the same corrupted frame. Imputation retains 649 records, whereas dropping leaves 617.",
    'The notebook displays missing counts for every column before simulation, after simulation and after each treatment. A comparison table records dimensions and the mean/sample standard deviation of absences. Both methods remove all missing cells, but they answer different practical questions: imputation retains the original population with constructed values, while dropping analyses a smaller set of fully observed records.'
)
R.table(['Variant','Rows','Columns','Missing','Mean absences','Sample SD'],[
    [r['variant'],r['rows'],r['columns'],r['missing_cells'],f"{r['absences_mean']:.4f}",f"{r['absences_std']:.4f}"] for r in S['missing_comparison']
], '2.3.1','Dimensions and absence summaries for each treatment',[4.4,1.5,1.8,1.8,3.1,2.9])
paras(
    'Mean imputation can insert a fractional value into an originally integer count. It also concentrates the replacement observations at one point and understates dispersion relative to the observed-data distribution. Row removal preserves the retained observed measurements but loses 32 records. Neither operation is presented as an unconditional best method.',
    'Because the complete original measurements are available, Parts II–IV use them for substantive findings. This avoids degrading known information solely to demonstrate preprocessing. The separate sensitivity comparison reports how the simulated treatments alter the relationship between absences and G3 and the membership of the illustrative two-condition filter.'
)
R.heading('2.4 Python structures and regular expressions',2)
paras(
    'The G3 Series is converted with tolist(). A Python list comprehension retains grades greater than or equal to 10. Its count is asserted to equal the equivalent Boolean mask. The threshold is an analytical choice used for this exercise; the performance bands later in the report are also explicitly analyst-defined.',
    'The regular expression G[123], applied using re.fullmatch, extracts G1, G2 and G3 from the column names. A full match rejects misleading strings such as G12 and G3_extra. This is a useful schema operation on the real data, rather than an invented dirty-data example or fabricated personal identification field.'
)

R.heading('CHAPTER 3. ANALYSIS AND TRANSFORMATION',new_page=True)
R.heading('3.1 Indexing and two-condition filtering',2)
paras(
    'The source frame has a RangeIndex from 0 to 648. The expression df.loc[:9] returns ten records because loc uses labels and includes the ending label. The iloc example selects the positional columns corresponding to school, age, G1, G2 and G3 on the original schema. The notebook displays both outputs and verifies their dimensions and column names.',
    f"Boolean indexing combines studytime ≥ 2 with absences ≤ 5 using a parenthesized conjunction. On the original data, {S['filter']['count']} of 649 records meet both conditions ({S['filter']['percent']:.2f}%). This is an illustrative subset, not a validated screening rule. Studytime ≥ 2 means membership in the 2–5 hour category or a higher category, not a continuous measurement of at least two coded hours."
)
R.heading('3.2 Derived features and their limits',2)
paras(
    'Three new columns satisfy the feature-engineering requirement. prior_grade_mean is inserted immediately before G3 with DataFrame.insert(); grade_progress subtracts G1 from G2; performance_band summarizes the final grade. Each formula and the insertion position are verified against the original measurements.',
    'The prior-grade mean avoids including the outcome G3. Nevertheless, it is a constructed combination of existing grades and is excluded from the ranking of original explanatory factors. Grade progress is an observed difference between assessments, not an isolated measure of teaching effectiveness. The performance band is directly derived from G3, so using it to explain G3 would be circular.'
)
R.equation('m','G1 + G2','2')
R.table(['New column','Definition','Analytical role'],[
    ['prior_grade_mean','(G1 + G2) / 2','Summarize earlier grades; inserted before G3'],
    ['grade_progress','G2 − G1','Describe change between two assessments'],
    ['performance_band','0–9; 10–14; 15–20','Describe the outcome; never rank as its explanation'],
    ['G3_minmax_school','Within-school min–max scale','Demonstrate groupwise transformation'],
], '3.2.1','Engineered columns and safeguards',[4.3,4.3,6.9])
R.heading('3.3 Correlation and association ranking',2)
prior=[r for r in S['rankings'] if r['family']=='Prior grades']
context=[r for r in S['rankings'] if r['family']!='Prior grades']
paras(
    'The required Pearson correlation matrix is calculated with corr() on the 16 original numerical columns. The complete matrix is retained in the notebook and CSV. Nominal labels are not converted into arbitrary numbers. Pearson summarizes linear relationships; ordinal scales and unusual observations require cautious interpretation.',
    'A Spearman table complements that matrix. It ranks original numerical and ordinal explanatory variables by the absolute value of their rank correlation with G3, while retaining the sign and sample size. G1 and G2 are shown separately from contextual variables. No derived feature, final-grade band or normalized outcome enters this ranking.',
    f"G2 has the strongest earlier-grade association with G3 (Spearman ρ = {prior[0]['spearman_rho']:.4f}; Pearson r = {S['pearson_G2_G3']:.4f}). G1 follows with ρ = {prior[1]['spearman_rho']:.4f}. These are earlier assessments of related achievement, so the result is not evidence that manipulating an earlier score causes a later score to change.",
    f"Among the examined original numerical/ordinal contextual factors, {context[0]['variable']} has the largest absolute Spearman coefficient (ρ = {context[0]['spearman_rho']:.4f}). The negative sign indicates that higher recorded failure categories are associated with lower ranks of G3. The upper code remains subject to the source documentation caveat."
)
R.table(['Original contextual variable','Spearman ρ','n'],[[r['variable'],f"{r['spearman_rho']:.4f}",r['n']] for r in context], '3.3.1','Contextual numerical/ordinal associations with G3',[8,4.5,3])
paras(
    'This ordering is descriptive and conditional on the variables examined. The table is not a feature-importance result from a multivariable model. Nominal-category mean differences are not placed on the same scale as correlation coefficients. Potential confounding, ties, questionnaire measurement error and the observational design limit causal interpretation.'
)
R.heading('3.4 Grouped grade statistics',2)
school_rows=[]
with (ROOT/'results/tables/school_grade_statistics.csv').open() as handle:
    rows=list(csv.reader(handle))
for row in rows[3:]:
    if len(row)>1:
        for i,grade in enumerate(['G1','G2','G3']):
            values=row[1+i*5:1+(i+1)*5]
            school_rows.append([row[0],grade,int(float(values[0])),f'{float(values[1]):.3f}',f'{float(values[2]):.3f}',f'{float(values[3]):.0f}',f'{float(values[4]):.0f}'])
# Two column-header rows and one index-name row precede the school data.
if len(school_rows)!=6:
    raise RuntimeError('School statistics CSV structure changed; inspect the generated artifact.')
R.table(['School','Grade','n','Mean','SD','Min','Max'],school_rows,'3.4.1','Descriptive statistics for all three grades by school',[2.1,2,1.7,2.8,2.7,2.1,2.1])
paras(
    f"For G3, GP has mean {S['school_g3'][0]['mean']:.3f} and median {S['school_g3'][0]['median']:.0f}; MS has mean {S['school_g3'][1]['mean']:.3f} and median {S['school_g3'][1]['median']:.0f}. GP contributes more records, while MS has greater observed dispersion. These differences describe the sampled groups and cannot be attributed to school quality without controlling for other differences.",
    'Grouping by studytime and failures addresses the relationship between one recorded factor and the outcome. Counts are essential because the higher study-time and failure categories are much smaller. The higher-education aspiration variable provides a nominal comparison; its group means are interpreted in their original units rather than compared numerically with Spearman coefficients.'
)
R.table(['Study-time code','Weekly interval','n','G3 mean','G3 SD'],[
    [r['studytime'],label,r['count'],f"{r['mean']:.3f}",f"{r['std']:.3f}"] for r,label in zip(S['group_stats']['studytime'],['<2 hours','2–5 hours','5–10 hours','>10 hours'])
], '3.4.2','Final-grade summaries by study-time category',[3.1,4.3,2,3,3.1])
R.table(['Recorded failures code','n','G3 mean','G3 SD'],[
    [r['failures'],r['count'],f"{r['mean']:.3f}",f"{r['std']:.3f}"] for r in S['group_stats']['failures']
], '3.4.3','Final-grade summaries by recorded failure category',[6,2.5,3.5,3.5])
R.heading('3.5 Within-school normalization',2)
paras(
    'The groupwise transform calculates each school minimum and maximum and broadcasts them back to the original row index. Dividing the centred grade by the school range gives a value between 0 and 1. If a group is constant, its transformed values are defined as 0 to avoid division by zero. The real groups and a deliberately constant test fixture are both checked.',
    'This transformation demonstrates relative scaling, not an improvement to the outcome. A normalized value locates a record within its school range; the general method does not make equal values from different schools equal in absolute achievement. The transformed column is never used as an explanatory factor for G3.'
)
R.equation('z','G3 − min(G3)','max(G3) − min(G3)')
R.heading('3.6 Sensitivity to simulated missingness',2)
R.table(['Variant','n','Absences SD','ρ with G3','Filter n'],[
    [r['variant'],r['n'],f"{r['absences_std']:.4f}",f"{r['absences_G3_spearman']:.4f}",r['filter_count']] for r in S['sensitivity']
], '3.6.1','Targeted sensitivity comparison using the same missing mask',[4.7,1.8,3.1,3.1,2.8])
paras(
    'The sensitivity comparison shows that treatment can change both dispersion and subset membership. Imputed observations fall below the absence threshold because their shared replacement mean is below five; this changes the Boolean-filter count. The dropna percentage uses 617 as its denominator and should not be compared as if it retained all original records.',
    'The sign of the absences association is stable in this particular simulation, but its magnitude changes. This single seeded MCAR experiment does not establish stability under other masks or missingness mechanisms. The complete original measurements therefore remain the appropriate basis for the main findings.'
)

R.heading('CHAPTER 4. DATA VISUALIZATION',new_page=True)
R.heading('4.1 Final-grade distribution',2)
stats=S['grade_statistics']['G3']
paras(f"Figure 4.1.1 uses Seaborn histplot with one-point-wide bins centred on integer G3 values. The mean final grade is {stats['mean']:.3f}, the median is {stats['median']:.0f}, and the sample standard deviation is {stats['std']:.3f}. The full 0–20 axis avoids hiding empty theoretical-score intervals.")
R.figure('01_grade_histogram.png','4.1.1','Distribution of final grades in the Portuguese-language dataset')
paras('The histogram concentrates most observations in the middle of the score range. A separate mass at zero should not be silently removed or reclassified as missing: zero is a recorded outcome. The distribution motivates reporting the median and spread alongside the mean, rather than describing achievement with a single average.')
R.heading('4.2 Distribution by school',2,new_page=True)
paras('Figure 4.2.1 uses Seaborn boxplot to compare the two school distributions. The median, interquartile range, 1.5-IQR whiskers and individually marked extreme observations summarize complementary aspects of the data. Sample sizes are included in the category labels.')
R.figure('02_school_boxplot.png','4.2.1','Final-grade distributions for GP and MS')
paras('GP has a higher observed central grade and a narrower distribution, while MS shows a lower centre and greater dispersion. The unequal group sizes and unadjusted composition prevent a causal comparison of school performance. Outlier markers identify positions relative to the plotting rule; they are not an automatic instruction to delete observations.')
R.heading('4.3 Comparison by study-time category',2,new_page=True)
paras('Figure 4.3.1 uses a Seaborn bar plot of mean G3 across the four original study-time categories. Each bar is accompanied by its mean and group size. The intervals are labelled in hours rather than showing the numeric codes as if they were measured quantities.')
R.figure('03_studytime_bar.png','4.3.1','Mean final grade by weekly study-time category')
paras('The pattern is consistent with a positive association between higher study-time categories and final grades, but the groups differ markedly in size. Study habits, prior achievement and other factors may be related. The figure omits an unexplained confidence band and does not claim that moving a student between categories would reproduce the observed mean difference.')
R.heading('4.4 Performance-band proportions',2,new_page=True)
paras('Figure 4.4.1 uses the performance_band column created in Part II. The bands cover every grade from 0 to 20 exactly once. Labels contain counts and percentages, and distinct academic colors (red for 0–9, blue for 10–14, and green for 15–20) provide clear visual encoding of student achievement levels.')
R.figure('04_performance_pie.png','4.4.1','Proportions of the analyst-defined final-grade bands')
bands=S['performance_distribution']
paras(f"The 0–9 band contains {bands['0–9']['count']} records ({bands['0–9']['percent']:.2f}%), the 10–14 band {bands['10–14']['count']} ({bands['10–14']['percent']:.2f}%), and the 15–20 band {bands['15–20']['count']} ({bands['15–20']['percent']:.2f}%). These are descriptive intervals selected for this report, not an official Portuguese grading classification. The proportions sum to 100% before display rounding.")
R.heading('4.5 Earlier and final grades',2,new_page=True)
paras('Figure 4.5.1 combines a Seaborn scatter plot with a descriptive linear regression line. Integer scores produce many overlapping coordinates; transparency reveals concentration without changing observations through jitter. The confidence ribbon is disabled because no interval method is being interpreted.')
R.figure('05_g2_g3_scatter.png','4.5.1','Relationship between second-period and final grades')
paras(f"The strong positive pattern agrees with Pearson r = {S['pearson_G2_G3']:.4f} and Spearman ρ = {prior[0]['spearman_rho']:.4f}. A line is a compact description, not a claim that the relation is exact at every score or a validated forecasting model. G2 precedes G3 and assesses related achievement, which is central to interpretation.")
R.heading('4.6 Matplotlib core demonstration',2,new_page=True)
paras('Figure 4.6.1 is constructed with Matplotlib alone from ordinary Python lists and a dictionary. Its values are the earlier-grade means for the first five source records. Pandas supplies the values before conversion, but no Pandas plotting or Seaborn API creates this graph.')
R.figure('06_matplotlib_core.png','4.6.1','Earlier-grade averages for the first five source records')
paras('The x-axis labels are one-based source-row positions, not personal identities. This five-record illustration demonstrates basic data structures and plotting mechanics; it is not a random sample or an estimate of a population trend. The complete dataset, rather than these five records, supports the report conclusions.')

R.heading('CHAPTER 5. SYNTHESIS AND CONCLUSION',new_page=True)
R.heading('5.1 Comprehensive reporting function',2)
paras(
    'The Python function comprehensive_report(df) accepts its dataset explicitly, displays a readable summary, and returns a JSON-safe dictionary. It reports total records and columns, missing cells per column, grade summaries, and counts/percentages for each performance band. The return value is extended by the notebook with the already-computed analysis artifacts used by the document builders.',
    f"For the primary analysis there are {S['total_records']} records and {S['total_columns']} columns after feature engineering and normalization. The four added columns are distinguishable from the original 33. Performance-band counts sum to the record count, percentages sum to 100 before rounding, and missing counts remain zero. Tests also cover an empty input with the correct schema and zero-count bands."
)
R.table(['Grade','n','Mean','Sample SD','Median','Min','Max'],[
    [grade,int(s['count']),f"{s['mean']:.3f}",f"{s['std']:.3f}",f"{s['median']:.0f}",f"{s['min']:.0f}",f"{s['max']:.0f}"] for grade,s in S['grade_statistics'].items()
], '5.1.1','Summary statistics of the three assessment grades',[2,1.8,2.6,2.9,2.4,1.9,1.9])
R.heading('5.2 Highest and lowest records',2)
paras(
    'The notebook displays all original and derived columns for exactly ten highest and ten lowest records by G3. A stable sort preserves original source order within tied grades. Therefore the displayed ten are a reproducible selection, not a claim that an omitted record with the same grade is worse or better.',
    'Tables 5.2.1 and 5.2.2 show readable selected columns. The complete 37-column evidence is retained in the notebook and the top_10_all_columns.csv and bottom_10_all_columns.csv exports. Source indexes in these tables are zero-based; the five-record chart uses explicitly labelled one-based row positions.'
)
for label,key,number in [('Highest','top_ten','5.2.1'),('Lowest','bottom_ten','5.2.2')]:
    R.table(['Source index','School','G1','G2','G3','Band'],[[r['source_index'],r['school'],r['G1'],r['G2'],r['G3'],r['performance_band']] for r in S[key]],number,f'{label} ten final-grade records: selected columns',[3.2,2.4,2,2,2,3.9])
R.heading('5.3 Findings and practical interpretation',2)
paras(
    f"Three findings summarize the evidence. The final-grade distribution has mean {stats['mean']:.3f} and median {stats['median']:.0f}; the 10–14 band contains the largest share ({bands['10–14']['percent']:.2f}%). Earlier G2 grades have a strong positive association with final grades. Among the original numerical/ordinal contextual factors examined, the recorded failures category has the strongest rank association with G3 (ρ = {context[0]['spearman_rho']:.4f}).",
    'The association ranking and grouped means justify closer examination of records with a history of difficulty, but they do not establish an effective intervention. A practical next study could investigate support needs using prospective measurements and explicit controls for prior achievement. No expected benefit, causal effect size or success rate is asserted from the present observations.',
    'The missingness exercise demonstrates why treatment decisions should be reported rather than hidden. Both treatments eliminate NaN values, yet they change the distribution or population being summarized. Retaining the known original measurements for the substantive analysis separates a teaching demonstration from the evidence used for conclusions.'
)
R.heading('5.4 Limitations',2)
paras(
    'The observations come from two Portuguese schools and one subject; they are not a representative sample of TDTU students or all educational settings. Questionnaire variables may involve measurement error. Ordinal codes, the ambiguity in the upper failures code, unequal group sizes and tied scores constrain interpretation.',
    'The analysis does not adjust for confounding or estimate multivariable effects. Correlation and group comparisons do not imply causation. The single missing-data simulation uses an MCAR mechanism and a fixed mask, and it cannot characterize all missingness patterns. The absence of a prediction model is deliberate because the assignment requires exploratory analysis and transformation rather than model evaluation.',
    'Reproducibility checks confirm that the documented operations execute and reproduce their outputs. They do not validate every source measurement or replace subject-matter judgement. These limits should be stated during the presentation, especially when discussing the apparently strongest factor.'
)
R.heading('5.5 Proposed presentation responsibilities',2)
paras('The following allocation is a proposal for preparing and rehearsing the presentation. It is not a statement of completed work or an assessment of individual contribution. Actual contribution percentages have not been supplied and are intentionally not fabricated.')
R.table(['Member','Proposed responsibility'],[
    ['NGO LAM TIEN — 524H0195','Introduction, dataset and Part I'],
    ['NGUYEN CONG THANH — 524H0191','Part II: features, groups and transformation'],
    ['NGUYEN LE NHAT TAM — 524H0186','Part III: charts and interpretation'],
    ['NGO BINH MINH — 524H0169','Part IV, code demonstration and conclusion'],
], '5.5.1','Proposed presentation allocation, not a contribution record',[7.3,8.2])

R.heading('REFERENCES',1,new_page=True)
R.doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
references=[
    ('Cortez, P. (2008). Student Performance [Dataset]. UCI Machine Learning Repository. CC BY 4.0. ', 'https://doi.org/10.24432/C5TG7T'),
    ('Pandas development team. (n.d.). Pandas documentation: indexing, missing data, groupby and correlation. Software version used: 2.2.3. ', 'https://pandas.pydata.org/docs/'),
    ('Seaborn developers. (n.d.). Seaborn API documentation: histplot, boxplot, barplot and regplot. Software version used: 0.13.2. ', 'https://seaborn.pydata.org/api.html'),
    ('Matplotlib development team. (n.d.). Matplotlib documentation: axes, bar and pie charts. Software version used: 3.10.3. ', 'https://matplotlib.org/stable/'),
    ('NumPy developers. (n.d.). NumPy documentation: random number generators and numerical arrays. Software version used: 2.2.6. ', 'https://numpy.org/doc/'),
]
for text,url in references:
    p=R.para(text,first=False);p.paragraph_format.left_indent=Cm(1.27);p.paragraph_format.first_line_indent=Cm(-1.27);p.paragraph_format.space_after=Pt(12)
    hyperlink(p,url)

R.finalize_indexes()
output=ROOT/'.build/report_draft.docx';R.doc.save(output)
(ROOT/'.build/report_index.json').write_text(json.dumps(R.entries+R.figures+R.tables,ensure_ascii=False,indent=2)+'\n')
print(output)
