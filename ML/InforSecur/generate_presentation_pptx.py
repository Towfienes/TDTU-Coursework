import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from playwright.sync_api import sync_playwright

# -------------------------------------------------------------
# Color Palette Constants
# -------------------------------------------------------------
CLR_BG = RGBColor(0x0B, 0x12, 0x20)          # Deep navy / near-black #0B1220
CLR_CARD = RGBColor(0x13, 0x1C, 0x2E)        # Dark graphite surface #131C2E
CLR_CARD_BORDER = RGBColor(0x1E, 0x29, 0x3B) # Subtle card border #1E293B
CLR_GREEN = RGBColor(0x3D, 0xDC, 0x84)       # Android green #3DDC84
CLR_CYAN = RGBColor(0x38, 0xBD, 0xF8)        # Cyan / sky blue #38BDF8
CLR_TEXT_WHITE = RGBColor(0xF8, 0xFA, 0xFC)  # Crisp near-white #F8FAFC
CLR_TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)  # Slate gray #94A3B8
CLR_TEXT_DIM = RGBColor(0x64, 0x74, 0x8B)    # Dim slate #64748B
CLR_RED = RGBColor(0xF8, 0x71, 0x71)         # Soft coral red #F87171
CLR_AMBER = RGBColor(0xFB, 0xBF, 0x24)       # Amber warning #FBBF24
CLR_CODE_BG = RGBColor(0x08, 0x0D, 0x18)     # Deep dark code box #080D18

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# -------------------------------------------------------------
# Helper Functions for PPTX Generation
# -------------------------------------------------------------
def set_slide_background(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = CLR_BG
    bg_shape.line.fill.background()
    # Send background to back by re-inserting or leaving as first shape
    return bg_shape

def add_header(slide, tag_text, title_text, subtitle_text=""):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_tag = tf.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = CLR_GREEN
    p_tag.space_after = Pt(2)

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = CLR_TEXT_WHITE
    p_title.space_after = Pt(2)

    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = CLR_TEXT_MUTED

def add_card(slide, left, top, width, height, fill_color=CLR_CARD, border_color=CLR_CARD_BORDER, border_width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape

def add_code_box(slide, left, top, width, height, code_text):
    card = add_card(slide, left, top, width, height, fill_color=CLR_CODE_BG, border_color=CLR_CYAN, border_width=1.0)
    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = code_text
    p.font.name = "Consolas"
    p.font.size = Pt(11)
    p.font.color.rgb = CLR_CYAN
    return card

def set_speaker_notes(slide, notes_text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = notes_text

# -------------------------------------------------------------
# Main Presentation Builder
# -------------------------------------------------------------
def build_pptx():
    prs = pptx.Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]

    # =========================================================
    # SLIDE 1: Title Slide
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Main Card container
    add_card(s1, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9), fill_color=CLR_CARD, border_color=CLR_CARD_BORDER, border_width=1.5)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(1.3), Inches(1.5), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CLR_GREEN
    bar.line.fill.background()

    # Title text box
    tb = s1.shapes.add_textbox(Inches(1.6), Inches(1.55), Inches(10.133), Inches(2.2))
